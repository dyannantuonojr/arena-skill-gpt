"""Bounded, data-only extraction in a short-lived subprocess. No OCR or execution."""
import csv
import io
from pathlib import Path
import zipfile

MAX_CHARS=500000
MAX_SECTIONS=50000


class DocumentError(ValueError): pass


def extract(path,extension):
    sections=[];warnings=[];chars=0
    def add(locator,text):
        nonlocal chars
        text=str(text)
        if not text.strip():return
        chars+=len(text)+len(locator)
        if chars>MAX_CHARS or len(sections)>=MAX_SECTIONS:
            raise DocumentError('Document exceeds extraction limits (500,000 characters or 50,000 populated locations). Split it into smaller files.')
        sections.append({'locator':locator,'text':text})
    if extension in ('.xlsx','.docx'):
        try:
            with zipfile.ZipFile(path) as archive:
                members=archive.infolist()
                if len(members)>10000 or sum(i.file_size for i in members)>50*1024*1024:
                    raise DocumentError('The document archive is too large when expanded.')
                if any(i.flag_bits&1 for i in members):raise DocumentError('Encrypted documents are not supported.')
        except zipfile.BadZipFile as exc:raise DocumentError('This is not a valid XLSX or DOCX file.') from exc
    if extension=='.xlsx':
        from openpyxl import load_workbook
        with open(path,'rb') as src:raw=src.read()
        formulas=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        cached=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        try:
            if len(formulas.worksheets)>100:raise DocumentError('Workbook exceeds 100 sheets.')
            warnings.append('Formulas are not executed. Cached formula values may be missing or stale; formatting, charts and images are not recalculated or analyzed.')
            for sheet,values in zip(formulas.worksheets,cached.worksheets):
                if sheet.max_row and sheet.max_column and (sheet.max_row>100000 or sheet.max_column>1000 or sheet.max_row*sheet.max_column>1000000):
                    raise DocumentError('Worksheet dimensions exceed extraction limits. Remove unused formatted rows/columns or split the workbook.')
                for row,crow in zip(sheet.iter_rows(),values.iter_rows()):
                    for cell,value in zip(row,crow):
                        if cell.value is None:continue
                        text=str(getattr(cell.value,'text',cell.value))
                        if cell.data_type=='f':text='formula: '+text+'; cached value: '+('unavailable' if value.value is None else str(value.value))
                        elif cell.number_format!='General':text+='; number format: '+str(cell.number_format)
                        add(sheet.title+'!'+cell.coordinate,text)
        finally:formulas.close();cached.close()
    elif extension=='.csv':
        def get_column_letter(number):
            name=''
            while number:
                number,remainder=divmod(number-1,26);name=chr(65+remainder)+name
            return name
        raw=Path(path).read_bytes()
        try:text=raw.decode('utf-16' if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig')
        except UnicodeDecodeError:
            text=raw.decode('cp1252');warnings.append('CSV encoding assumed Windows-1252 after UTF-8 decoding failed.')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t|')
        except csv.Error:dialect=csv.excel
        for rownum,row in enumerate(csv.reader(io.StringIO(text),dialect),1):
            if rownum>100000 or len(row)>1000:raise DocumentError('CSV dimensions exceed extraction limits.')
            for col,value in enumerate(row,1):add('CSV!'+get_column_letter(col)+str(rownum),value)
        warnings.append('CSV cells are treated as supplied text; formula-like values are not executed.')
    elif extension=='.pdf':
        from pypdf import PdfReader,overwrite_configuration
        from pypdf.errors import LimitReachedError
        cap=8*1024*1024
        overwrite_configuration(maximum_declared_stream_length=cap,array_based_stream_maximum_output_length=cap,
            zlib_maximum_output_length=cap,lzw_maximum_output_length=cap,run_length_maximum_output_length=cap,
            brotli_maximum_output_length=cap,jbig2_maximum_output_length=cap,image_maximum_buffer_size=cap,
            stream_decoding_work_maximum_length=4*cap,page_tree_maximum_entries=1000,jbig2dec_binary='')
        try:
            reader=PdfReader(path)
            if reader.is_encrypted:raise DocumentError('Encrypted PDFs are not supported. Upload an unlocked copy.')
            if len(reader.pages)>200:raise DocumentError('PDF exceeds 200 pages. Split it into smaller files.')
            for number,page in enumerate(reader.pages,1):
                text=page.extract_text() or ''
                if text.strip():add('page '+str(number),text)
                else:warnings.append('Page '+str(number)+' has no extractable text; image/OCR content was not read.')
        except LimitReachedError as exc:raise DocumentError('PDF exceeds bounded stream/processing limits. Split or simplify it before uploading.') from exc
        warnings.append('PDF text extraction may not preserve visual table layout. Images and scanned text are not OCR-processed.')
    elif extension=='.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        document=Document(path);paragraph=table=0
        for block in document.element.body:
            if block.tag.endswith('}p'):
                paragraph+=1;add('paragraph '+str(paragraph),Paragraph(block,document).text)
            elif block.tag.endswith('}tbl'):
                table+=1
                for rownum,row in enumerate(Table(block,document).rows,1):
                    for col,cell in enumerate(row.cells,1):add('table %d row %d cell %d'%(table,rownum,col),cell.text)
        warnings.append('DOCX body paragraphs and tables extracted. Images, text boxes, headers/footers, footnotes, nested tables, comments and revision history are not analyzed.')
    else:
        try:text=Path(path).read_text(encoding='utf-8-sig')
        except UnicodeDecodeError as exc:raise DocumentError('Text and Markdown files must use UTF-8 encoding.') from exc
        for number,line in enumerate(text.splitlines(),1):add('line '+str(number),line)
    if not sections:raise DocumentError('The file has no readable text. For a scanned PDF, provide an OCR/text version.')
    return {'version':1,'sections':sections,'warnings':warnings}
