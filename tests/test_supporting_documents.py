"""Synthetic supporting-document fixtures, no model workers."""
import base64
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'skills/arena'))
from supporting_documents import DocumentStore, snapshot_documents, verify_run_documents


class DocumentsTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name); self.store=DocumentStore(self.root/'uploads')

    def upload(self,data=b'Item,Amount\nWidgets,42\n',name='budget.csv',key='a'*32):
        return self.store.upload(key,name,len(data),hashlib.sha256(data).hexdigest(),0,base64.b64encode(data).decode())

    def test_csv_upload_snapshot_and_source_location(self):
        result=self.upload()
        self.assertTrue(result['ready'])
        self.assertEqual(result['name'],'budget.csv')
        preview=self.store.preview(result['id'],'CSV!B2')
        self.assertIn('42',preview['text'])
        sources=self.store.resolve([result['id']])
        run=self.root/'run';run.mkdir()
        snapshot_documents(run,sources)
        manifest,hashes=verify_run_documents(run)
        self.assertEqual(manifest['documents'][0]['name'],'budget.csv')
        self.assertGreater(len(hashes),2)
        source=run/'supporting'/result['id']/'source.csv'
        source.write_bytes(b'changed')
        with self.assertRaises(ValueError):verify_run_documents(run)

    def test_chunk_retry_and_changed_chunk(self):
        data=b'hello world';digest=hashlib.sha256(data).hexdigest();key='b'*32
        args=(key,'notes.txt',len(data),digest,0,base64.b64encode(data[:5]).decode())
        self.assertEqual(self.store.upload(*args)['received'],5)
        self.assertEqual(self.store.upload(*args)['received'],5)
        with self.assertRaises(ValueError):self.store.upload(key,'notes.txt',len(data),digest,0,base64.b64encode(b'wrong').decode())
        result=self.store.upload(key,'notes.txt',len(data),digest,5,base64.b64encode(data[5:]).decode())
        self.assertTrue(result['ready'])
        self.assertTrue(self.store.upload(key,'notes.txt',len(data),digest,5,base64.b64encode(data[5:]).decode())['ready'])

    def test_invalid_path_type_size_hash_and_gap(self):
        for name in ('../evil.csv','folder\\evil.csv','macros.xlsm','program.exe'):
            with self.assertRaises(ValueError):self.upload(name=name)
        with self.assertRaises(ValueError):self.store.upload('c'*32,'x.txt',20*1024*1024,'0'*64,0,'eA==')
        with self.assertRaises(ValueError):self.store.upload('d'*32,'x.txt',1,'0'*64,0,'eA==')
        with self.assertRaises(ValueError):self.store.upload('e'*32,'x.txt',2,hashlib.sha256(b'xx').hexdigest(),1,'eA==')

    def test_snapshot_is_independent_of_staging(self):
        result=self.upload();run=self.root/'run';run.mkdir()
        snapshot_documents(run,self.store.resolve([result['id']]))
        (self.root/'uploads'/result['id']/'source.csv').write_bytes(b'changed staging')
        self.assertEqual(verify_run_documents(run)[0]['documents'][0]['name'],'budget.csv')
        with self.assertRaises(ValueError):self.store.resolve([result['id']])

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'),'Optional document parsers not installed')
    def test_xlsx_preserves_cell_formula_and_missing_cache(self):
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.title='Budget';sheet['A1']='Amount';sheet['B2']=42;sheet['C2']='=B2*2'
        data=io.BytesIO();book.save(data)
        result=self.upload(data.getvalue(),'budget.xlsx')
        preview=self.store.preview(result['id'],'Budget!B2:C2')
        self.assertIn('42',preview['text']);self.assertIn('=B2*2',preview['text'])
        self.assertIn('unavailable',preview['text'])

    @unittest.skipUnless(importlib.util.find_spec('docx'),'Optional document parsers not installed')
    def test_docx_paragraph_and_table_locators(self):
        from docx import Document
        doc=Document();doc.add_paragraph('A supporting statement');doc.add_table(1,1).cell(0,0).text='Table evidence'
        data=io.BytesIO();doc.save(data)
        result=self.upload(data.getvalue(),'support.docx')
        self.assertIn('supporting statement',self.store.preview(result['id'],'paragraph 1')['text'])
        self.assertIn('Table evidence',self.store.preview(result['id'],'table 1 row 1 cell 1')['text'])

    @unittest.skipUnless(importlib.util.find_spec('pypdf'),'Optional PDF parser not installed')
    def test_pdf_page_locator_and_encryption_rejection(self):
        objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
        stream=b'BT /F1 12 Tf 72 720 Td (Supporting PDF evidence) Tj ET'
        objects.append(b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream')
        data=b'%PDF-1.4\n';offsets=[0]
        for number,body in enumerate(objects,1):
            offsets.append(len(data));data+=str(number).encode()+b' 0 obj\n'+body+b'\nendobj\n'
        xref=len(data);data+=b'xref\n0 6\n0000000000 65535 f \n'
        for offset in offsets[1:]:data+=('%010d 00000 n \n'%offset).encode()
        data+=b'trailer << /Size 6 /Root 1 0 R >>\nstartxref\n'+str(xref).encode()+b'\n%%EOF\n'
        result=self.upload(data,'evidence.pdf')
        self.assertIn('Supporting PDF evidence',self.store.preview(result['id'],'page 1')['text'])
        import zlib
        compressed=zlib.compress(b' '*(9*1024*1024)+stream)
        objects[-1]=b'<< /Length '+str(len(compressed)).encode()+b' /Filter /FlateDecode >>\nstream\n'+compressed+b'\nendstream'
        oversized=b'%PDF-1.4\n';positions=[]
        for number,body in enumerate(objects,1):
            positions.append(len(oversized));oversized+=str(number).encode()+b' 0 obj\n'+body+b'\nendobj\n'
        start=len(oversized);oversized+=b'xref\n0 6\n0000000000 65535 f \n'
        for position in positions:oversized+=('%010d 00000 n \n'%position).encode()
        oversized+=b'trailer << /Size 6 /Root 1 0 R >>\nstartxref\n'+str(start).encode()+b'\n%%EOF\n'
        with self.assertRaises(ValueError):self.upload(oversized,'expanded.pdf','c'*32)
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.encrypt('password');buffer=io.BytesIO();writer.write(buffer)
        with self.assertRaises(ValueError):self.upload(buffer.getvalue(),'locked.pdf','b'*32)

    def test_redirected_uploads_are_rejected(self):
        self.upload()
        meta=self.root/'uploads'/('a'*32)/'metadata.json'
        value=json.loads(meta.read_text(encoding='utf-8'))
        value['files']['../outside.txt']='0'*64
        meta.write_text(json.dumps(value),encoding='utf-8')
        with self.assertRaises(ValueError):self.store.resolve(['a'*32])

    def test_long_location_preview_shows_bounded_prefix(self):
        result=self.upload(b'A'*40001,'long.txt')
        preview=self.store.preview(result['id'],'line 1')
        self.assertTrue(preview['text'])
        self.assertTrue(preview['partial'])
        self.assertIn('preview truncated',preview['text'])

    def test_unknown_document_and_locator_rejected(self):
        result=self.upload()
        with self.assertRaises(ValueError):self.store.preview(result['id'],'CSV!Z9999')
        with self.assertRaises(ValueError):self.store.resolve(['f'*32])

if __name__=='__main__':unittest.main()
